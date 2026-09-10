from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.db.models import Sum
from rest_framework import serializers

from compras.models import PedidoCompraItem
from compras.serializers import PedidoCompraItemSerializer
from produto.models import Cor, Grade, Pack, PackItem, Produto, ProdutoDetalhe, ProdutoFornecedor


REQUIRED_COLUMNS = {
    "Codigo_Produto_Fornecedor",
    "Produto",
    "Grade",
    "Cor",
    "Pack",
    "Nr_Packs",
    "Preco_Unitario",
    "Desconto",
    "Observacoes",
}

EXTRA_COLUMNS = {
    "Ordem",
    "Qtde_Por_Pack",
    "Cores_Expandidas",
    "Qtde_Total_Prevista",
    "Total_Previsto",
}


class PedidoCompraImportacaoError(Exception):
    pass


@dataclass
class ImportRow:
    source_line: int
    data: dict


def _norm_header(value) -> str:
    return str(value or "").strip()


def _norm_text(value) -> str:
    return str(value or "").strip()


def _norm_key(value) -> str:
    return _norm_text(value).casefold()


def _decimal(value, field_name: str, errors: list[str], line: int, default=None):
    if value in (None, ""):
        if default is not None:
            return default
        errors.append(f"Linha {line}: {field_name} inválido.")
        return None
    try:
        parsed = Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, ValueError):
        errors.append(f"Linha {line}: {field_name} inválido.")
        return None
    if parsed < 0:
        errors.append(f"Linha {line}: {field_name} deve ser maior ou igual a zero.")
        return None
    return parsed


def _int_packs(value, errors: list[str], line: int):
    if value in (None, ""):
        errors.append(f"Linha {line}: Nº de packs é obrigatório.")
        return None
    try:
        parsed = Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, ValueError):
        errors.append(f"Linha {line}: Nº de packs deve ser inteiro maior ou igual a 1.")
        return None
    if parsed != parsed.to_integral_value() or parsed < 1:
        errors.append(f"Linha {line}: Nº de packs deve ser inteiro maior ou igual a 1.")
        return None
    return int(parsed)


def read_xlsx_rows(uploaded_file) -> list[ImportRow]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise PedidoCompraImportacaoError("Biblioteca openpyxl não instalada no backend.") from exc

    try:
        content = uploaded_file.read()
        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise PedidoCompraImportacaoError("Arquivo XLSX inválido ou corrompido.") from exc

    if "Itens" not in workbook.sheetnames:
        raise PedidoCompraImportacaoError('A planilha deve possuir uma aba chamada "Itens".')

    sheet = workbook["Itens"]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise PedidoCompraImportacaoError("A aba Itens não possui cabeçalho.")

    headers = [_norm_header(cell) for cell in header_row]
    header_map = {header.casefold(): idx for idx, header in enumerate(headers) if header}
    missing = sorted(col for col in REQUIRED_COLUMNS if col.casefold() not in header_map)
    if missing:
        raise PedidoCompraImportacaoError(f"Colunas obrigatórias ausentes: {', '.join(missing)}.")

    allowed = {c.casefold() for c in REQUIRED_COLUMNS | EXTRA_COLUMNS}
    rows: list[ImportRow] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        data = {}
        for header, value in zip(headers, values):
            if header and header.casefold() in allowed:
                data[header] = value
        rows.append(ImportRow(source_line=row_number, data=data))
    return rows


class PedidoCompraImportacaoService:
    def __init__(self, pedido, request=None):
        self.pedido = pedido
        self.request = request
        self.empresa = pedido.empresa
        self.errors: list[str] = []

    def preview_from_file(self, uploaded_file):
        rows = read_xlsx_rows(uploaded_file)
        return self.preview(rows)

    def preview_from_payload(self, rows_payload):
        rows = [ImportRow(source_line=int(row.get("linha_planilha") or row.get("source_line") or 0), data=row) for row in rows_payload]
        return self.preview(rows)

    def preview(self, rows: list[ImportRow]):
        self.errors = []
        if self.pedido.status != "AB":
            self.errors.append("Somente pedidos em aberto (AB) permitem importar itens.")
        if self.pedido.cotacao_origem_id:
            self.errors.append("Pedido originado de cotação aprovada não permite alteração comercial.")

        preview_rows = []
        for row in rows:
            preview_rows.extend(self._expand_row(row))

        self._validate_duplicates(preview_rows)
        valid = not self.errors
        return {
            "valid": valid,
            "errors": self.errors,
            "total_linhas_planilha": len(rows),
            "total_linhas_preview": len(preview_rows),
            "total_quantidade": str(sum((r["quantidade_calculada"] for r in preview_rows), Decimal("0"))),
            "total_valor": str(sum((r["total"] for r in preview_rows), Decimal("0")).quantize(Decimal("0.01"))),
            "linhas": preview_rows,
        }

    def confirm(self, linhas):
        rows = [ImportRow(source_line=int(row.get("linha_planilha") or 0), data=row) for row in linhas]
        preview = self.preview(rows)
        if not preview["valid"]:
            raise serializers.ValidationError({"errors": preview["errors"], "linhas": preview["linhas"]})
        with transaction.atomic():
            pedido = type(self.pedido).objects.select_for_update().get(pk=self.pedido.pk)
            self.pedido = pedido
            self.empresa = pedido.empresa
            if pedido.status != "AB" or pedido.cotacao_origem_id:
                raise serializers.ValidationError({"detail": "Pedido não permite importação neste momento."})
            created = []
            for row in preview["linhas"]:
                serializer = PedidoCompraItemSerializer(data={
                    "pedido": pedido.pk,
                    "produto": row["produto_id"],
                    "cor": row["cor_id"],
                    "pack": row["pack_id"],
                    "n_packs": row["n_packs"],
                    "preco_unit": row["preco_unitario"],
                    "desconto_valor": row["desconto"],
                    "observacoes": row["observacoes"] or None,
                }, context={"request": self.request})
                serializer.is_valid(raise_exception=True)
                created.append(serializer.save())
            pedido.refresh_from_db()
            pedido.recomputa_totais()
            pedido.save(update_fields=["tipo", "total_itens", "total_desconto", "frete", "outras_despesas", "total_pedido"])
        return {
            "pedido_id": pedido.pk,
            "itens_criados": len(created),
            "referencias": len({item.produto_id for item in created}),
            "quantidade_total": str(sum((item.qtd for item in created), Decimal("0"))),
            "total_importado": str(sum((item.total_item for item in created), Decimal("0")).quantize(Decimal("0.01"))),
            "pedido": pedido,
        }

    def _expand_row(self, row: ImportRow):
        data = row.data
        line = row.source_line
        produto = self._resolve_produto(data, line)
        n_packs = _int_packs(self._get(data, "Nr_Packs"), self.errors, line)
        preco = _decimal(self._get(data, "Preco_Unitario"), "preço unitário", self.errors, line)
        desconto = _decimal(self._get(data, "Desconto"), "desconto", self.errors, line, Decimal("0"))
        observacoes = _norm_text(self._get(data, "Observacoes"))
        pack = None
        cores = []
        if produto:
            self._validate_grade(produto, self._get(data, "Grade"), line)
            pack = self._resolve_pack(produto, self._get(data, "Pack"), line)
            cores = self._resolve_cores(produto, self._get(data, "Cor"), line)
        if not produto or n_packs is None or preco is None or desconto is None:
            return []
        if not pack or not cores:
            return []
        qtd_pack = PackItem.objects.filter(pack=pack).aggregate(total=Sum("qtd"))["total"] or 0
        quantidade = Decimal(int(qtd_pack) * int(n_packs))
        rows = []
        for index, cor in enumerate(cores, start=1):
            total = (quantidade * preco - desconto).quantize(Decimal("0.01"))
            rows.append({
                "linha_planilha": line,
                "origem": f"{line}.{index}" if len(cores) > 1 else str(line),
                "codigo_produto_fornecedor": _norm_text(self._get(data, "Codigo_Produto_Fornecedor")),
                "produto_id": produto.pk,
                "produto": produto.descricao,
                "produto_referencia": produto.referencia or "",
                "grade": getattr(produto.grade, "Descricao", "") or "",
                "cor_id": cor.pk,
                "cor": cor.Descricao,
                "pack_id": pack.pk,
                "pack": pack.nome or f"Pack {pack.pk}",
                "n_packs": n_packs,
                "quantidade_calculada": quantidade,
                "preco_unitario": preco,
                "desconto": desconto,
                "total": total,
                "observacoes": observacoes,
                "situacao": "OK",
            })
        return rows

    def _get(self, data, column):
        aliases = {
            "Codigo_Produto_Fornecedor": ["Codigo_Produto_Fornecedor", "codigo_produto_fornecedor"],
            "Produto": ["Produto", "produto", "produto_referencia"],
            "Grade": ["Grade", "grade"],
            "Cor": ["Cor", "cor"],
            "Pack": ["Pack", "pack"],
            "Nr_Packs": ["Nr_Packs", "n_packs"],
            "Preco_Unitario": ["Preco_Unitario", "preco_unitario"],
            "Desconto": ["Desconto", "desconto"],
            "Observacoes": ["Observacoes", "observacoes"],
        }
        for key, value in data.items():
            if _norm_key(key) in {alias.casefold() for alias in aliases.get(column, [column])}:
                return value
        return None

    def _resolve_produto(self, data, line):
        candidates = []
        codigo = _norm_text(self._get(data, "Codigo_Produto_Fornecedor"))
        if codigo:
            qs = ProdutoFornecedor.objects.select_related("produto").filter(
                empresa=self.empresa,
                fornecedor=self.pedido.fornecedor,
                ativo=True,
                codigo_vigente=ProdutoFornecedor.normalizar_codigo(codigo),
                produto__ativo=True,
                produto__tipo_produto="1",
            )
            candidates += [v.produto for v in qs]
            if not candidates:
                self.errors.append(f"Linha {line}: produto não encontrado para este fornecedor.")
        produto_ref = _norm_text(self._get(data, "Produto"))
        produto_by_ref = None
        if produto_ref:
            matches = list(Produto.objects.filter(empresa=self.empresa, ativo=True, tipo_produto="1").filter(
                Q_by_reference(produto_ref)
            )[:2])
            if len(matches) == 1:
                produto_by_ref = matches[0]
            elif len(matches) > 1 and not codigo:
                self.errors.append(f"Linha {line}: produto/referência Sysvar ambíguo.")
        if produto_by_ref:
            candidates.append(produto_by_ref)
        unique = {p.pk: p for p in candidates}
        if not unique:
            return None
        if len(unique) > 1:
            self.errors.append(f"Linha {line}: identificações de produto apontam para produtos diferentes.")
            return None
        produto = next(iter(unique.values()))
        if produto.tipo_produto != "1":
            self.errors.append(f"Linha {line}: produto pertence a outro tipo de compra.")
            return None
        if self.pedido.tipo and self.pedido.tipo != "1":
            self.errors.append(f"Linha {line}: pedido já possui tipo incompatível com Revenda.")
            return None
        return produto

    def _validate_grade(self, produto, grade_value, line):
        grade_text = _norm_text(grade_value)
        if not grade_text:
            return
        grade_real = getattr(produto.grade, "Descricao", "") or ""
        if grade_text.casefold() != grade_real.casefold():
            self.errors.append(f"Linha {line}: grade divergente da grade real do produto.")

    def _resolve_pack(self, produto, pack_value, line):
        pack_text = _norm_text(pack_value)
        if not pack_text:
            self.errors.append(f"Linha {line}: pack é obrigatório.")
            return None
        qs = Pack.objects.filter(empresa=self.empresa, ativo=True, grade=produto.grade).filter(nome__iexact=pack_text)
        pack = qs.first()
        if not pack:
            self.errors.append(f"Linha {line}: pack incompatível com a grade {getattr(produto.grade, 'Descricao', '')}.")
            return None
        if not PackItem.objects.filter(pack=pack).exists():
            self.errors.append(f"Linha {line}: pack não possui itens.")
            return None
        return pack

    def _resolve_cores(self, produto, cor_value, line):
        cor_text = _norm_text(cor_value)
        if not cor_text:
            self.errors.append(f"Linha {line}: cor é obrigatória.")
            return []
        cor_ids = ProdutoDetalhe.objects.filter(produto=produto).values_list("idcor_id", flat=True).distinct()
        sku_colors = Cor.objects.filter(pk__in=cor_ids).order_by("Descricao", "Idcor")
        if cor_text.casefold() == "todas":
            return list(sku_colors)
        matches = list(sku_colors.filter(Q_cor(cor_text))[:2])
        if len(matches) == 1:
            return matches
        self.errors.append(f"Linha {line}: cor \"{cor_text}\" não pertence ao produto ou é ambígua.")
        return []

    def _validate_duplicates(self, rows):
        seen = {}
        for row in rows:
            key = (row["produto_id"], row["cor_id"], row["pack_id"])
            if key in seen:
                self.errors.append(f"Linha {row['linha_planilha']}: duplicidade Produto + Cor + Pack com linha {seen[key]}.")
            else:
                seen[key] = row["linha_planilha"]
        for row in rows:
            exists = PedidoCompraItem.objects.filter(
                pedido=self.pedido,
                produto_id=row["produto_id"],
                cor_id=row["cor_id"],
                pack_id=row["pack_id"],
            ).exists()
            if exists:
                self.errors.append(f"Linha {row['linha_planilha']}: item já existe no pedido para Produto + Cor + Pack.")


def Q_by_reference(value):
    from django.db.models import Q

    return Q(referencia__iexact=value) | Q(descricao__iexact=value) | Q(descricao_reduzida__iexact=value)


def Q_cor(value):
    from django.db.models import Q

    return Q(Codigo__iexact=value) | Q(Descricao__iexact=value) | Q(Cor__iexact=value)
